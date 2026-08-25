#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""สร้าง Muhammadhisam-Pala-Portfolio-2026-editable.pptx

ต่างจาก make-pptx.py ตรงที่ไฟล์นี้วางข้อความเป็น text box จริง
เปิดใน PowerPoint / Keynote / Google Slides แล้วแก้ข้อความได้ทันที

ต้องมี:
    pip install python-pptx
    brew install imagemagick        # ใช้แปลง .webp เป็น .jpg

วิธีใช้:
    python3 make-pptx-editable.py
"""
import os
import re
import subprocess
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets", "projects")
BUILD = os.path.join(HERE, ".pptx-build")
OUT = os.path.join(HERE, "Muhammadhisam-Pala-Portfolio-2026-editable.pptx")

# ---------------------------------------------------------------- พื้นฐาน
# ออกแบบบนพิกัด 1920x1080 px แล้วแปลงเป็น EMU (สไลด์ 13.333 x 7.5 นิ้ว)
SLIDE_W_IN, SLIDE_H_IN = 13.333, 7.5
PX = int(round(SLIDE_W_IN * 914400 / 1920))          # EMU ต่อ 1 px
PT_PER_PX = SLIDE_W_IN * 72 / 1920                    # pt ต่อ 1 px

NAVY = RGBColor(0x07, 0x2A, 0x57)
ACCENT = RGBColor(0xE8, 0x55, 0x2F)
BODY = RGBColor(0x55, 0x60, 0x6D)
LABEL = RGBColor(0x97, 0xA2, 0xAE)
HAIR = RGBColor(0xD8, 0xDE, 0xE4)
WASH = RGBColor(0xF6, 0xF8, 0xF9)
PANEL = RGBColor(0xE4, 0xE7, 0xEA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

LATIN_FONT = "Arial"      # มีทุกเครื่อง
THAI_FONT = "Tahoma"      # มีทุกเครื่องและรองรับไทย
THAI_RE = re.compile(r"[฀-๿]")


def font_for(text):
    return THAI_FONT if THAI_RE.search(text) else LATIN_FONT


def emu(px):
    return Emu(int(round(px * PX)))


def pt(px):
    return Pt(px * PT_PER_PX)


# ---------------------------------------------------------------- ชิ้นส่วน
def rect(slide, x, y, w, h, fill=None, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(x), emu(y), emu(w), emu(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    return sh


def text(slide, x, y, w, h, blocks, size=18, color=BODY, bold=False,
         align=PP_ALIGN.LEFT, spacing=1.35, caps=False, letter=None,
         anchor=MSO_ANCHOR.TOP, fit=True):
    """blocks = str | [(text, {size,color,bold}), ...] | [[para], [para]]"""
    box = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if fit:
        # ให้ PowerPoint ย่อฟอนต์อัตโนมัติถ้าข้อความยาวเกินกรอบ
        # (ฟอนต์ไทย Tahoma กว้างกว่าฟอนต์ที่ใช้ในเว็บ ข้อความบางบล็อกจึงอาจล้น)
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    if isinstance(blocks, str):
        blocks = [[(blocks, {})]]
    elif blocks and isinstance(blocks[0], tuple):
        blocks = [blocks]

    for i, para in enumerate(blocks):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        for chunk, opt in para:
            shown = chunk.upper() if caps else chunk
            r = p.add_run()
            r.text = shown
            f = r.font
            f.name = font_for(shown)
            f.size = pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.color.rgb = opt.get("color", color)
            if letter or opt.get("letter"):
                # letter-spacing ไม่มีใน python-pptx ต้องเซ็ตผ่าน XML
                r.font._rPr.set("spc", str(int((opt.get("letter", letter)) * 100)))
    return box


def picture(slide, x, y, w, h, path, crop_top_bias=True):
    """วางรูปแบบ cover — ครอบเต็มกรอบโดยไม่ยืดผิดสัดส่วน"""
    pic = slide.shapes.add_picture(path, emu(x), emu(y))
    ai = pic.width / pic.height
    ab = float(w) / float(h)
    if ai > ab:                      # รูปกว้างเกิน -> ตัดซ้ายขวา
        keep = ab / ai
        pic.crop_left = pic.crop_right = (1 - keep) / 2
    else:                            # รูปสูงเกิน -> ตัดล่าง (ยึดขอบบน)
        keep = ai / ab
        if crop_top_bias:
            pic.crop_top, pic.crop_bottom = 0.0, 1 - keep
        else:
            pic.crop_top = pic.crop_bottom = (1 - keep) / 2
    pic.left, pic.top, pic.width, pic.height = emu(x), emu(y), emu(w), emu(h)
    return pic


def picture_fit(slide, x, y, w, h, path):
    """วางรูปแบบ contain — เห็นเต็มรูป จัดกึ่งกลางกรอบ"""
    pic = slide.shapes.add_picture(path, emu(x), emu(y))
    ai = pic.width / pic.height
    if ai > float(w) / float(h):
        nw, nh = w, w / ai
    else:
        nh, nw = h, h * ai
    pic.left, pic.top = emu(x + (w - nw) / 2), emu(y + (h - nh) / 2)
    pic.width, pic.height = emu(nw), emu(nh)
    return pic


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def chrome(slide, eyebrow, title=None, lede=None, rule="full", narrow_foot=False):
    """หัวสไลด์มาตรฐาน: eyebrow + เส้นคั่น + หัวเรื่อง + คำโปรย + footer"""
    text(slide, 80, 22, 1400, 30, eyebrow, size=15, color=LABEL, bold=True,
         caps=True, letter=2.1)
    if rule == "full":
        rect(slide, 80, 63, 1760, 1, fill=HAIR)
    elif rule == "half":
        rect(slide, 80, 62, 600, 4, fill=ACCENT)
    if title:
        text(slide, 80, 100, 1760 if rule == "full" else 570, 130, title,
             size=50, color=NAVY, bold=True, spacing=1.15, fit=False)
    if lede:
        text(slide, 80, 200, 1600, 90, lede, size=24, color=BODY, spacing=1.45)
    foot(slide, narrow=narrow_foot)


def foot(slide, narrow=False):
    w = 600 if narrow else 1760
    rect(slide, 80, 1020, w, 1, fill=HAIR)
    text(slide, 80, 1034, w, 30,
         [("Muhammadhisam Pala", {"bold": True, "color": NAVY}),
          ("   |   Software Developer Portfolio   |   2026", {})],
         size=16, color=LABEL)


def label(slide, x, y, w, s):
    text(slide, x, y, w, 24, s, size=14, color=ACCENT, bold=True, caps=True, letter=2.0)


def metrics(slide, x, y, items):
    """แถบตัวเลข impact"""
    for i, (num, cap) in enumerate(items):
        cx = x + i * 260
        rect(slide, cx, y, 48, 4, fill=ACCENT)
        text(slide, cx, y + 18, 240, 56, num, size=40, color=NAVY, bold=True, spacing=1.0, fit=False)
        text(slide, cx, y + 66, 240, 60, [[(l, {})] for l in cap], size=14,
             color=LABEL, spacing=1.3)


# ---------------------------------------------------------------- เนื้อหา
PROJECTS = [
    dict(key="erc", eyebrow="Aquila Studio · Energy Platform",
         title=["Centralized", "Energy Data Platform"],
         desc="ระบบรวบรวมและแสดงผลข้อมูลความต้องการใช้ไฟฟ้าและกำลังการผลิตจากหลายแหล่งไว้ที่เดียว "
              "พร้อม dashboard ที่ดูแยกรายภูมิภาค ประเภทเชื้อเพลิง และสัญญาซื้อขายไฟฟ้าได้",
         platform="Responsive Web Dashboard",
         work="Backend · REST API · Data Pipeline · Database Design",
         stack="NestJS · TypeScript · PostgreSQL · Docker",
         overview="ร่วมพัฒนา Energy Data Aggregation API ที่รวมข้อมูลไฟฟ้าจากหลายแหล่งเข้า dashboard เดียว "
                  "สำหรับดูภาพรวมและ monitoring ระดับประเทศ ผมดูแลฝั่ง backend ตั้งแต่ออกแบบโครงสร้างข้อมูล "
                  "เขียน API รวมข้อมูล จนถึง deploy บน cloud",
         cover="erc-01", cover_cap="แดชบอร์ดภาพรวมการใช้ไฟฟ้าทั้งประเทศ พร้อมแผนที่โรงไฟฟ้ารายจุด",
         gal_lede="หน้าจอเพิ่มเติมของระบบ — แดชบอร์ดรายภูมิภาค คลังข้อมูลย้อนหลัง และหน้านำเข้าข้อมูล",
         gallery=[("erc-02", "แดชบอร์ดกำลังผลิตแยกรายภูมิภาคและประเภทเชื้อเพลิง"),
                  ("erc-04", "คลังข้อมูลย้อนหลัง แยกตามภูมิภาคและประเภทรายงาน"),
                  ("erc-05", "หน้านำเข้าข้อมูล Supply / Demand / MA Plan")]),

    dict(key="pool", eyebrow="Aquila Studio · Gas Management",
         title=["Pool Manager", "Gas Pool Management"],
         desc="ระบบบริหารจัดการก๊าซธรรมชาติ คำนวณต้นทุนจากแหล่งจัดหา (LNG, Gulf, Myanmar) "
              "อัตราแลกเปลี่ยน และข้อมูล Shipper จากหลายระบบ",
         platform="Web Application (Internal)",
         work="Backend · Cost Allocation Engine · Audit Trail",
         stack="TypeScript · PostgreSQL · Docker",
         overview="รองรับ Cost Allocation & Pricing Workflow และเก็บประวัติให้ตรวจสอบย้อนหลังได้ "
                  "ทุกการคำนวณต้องอธิบายที่มาของตัวเลขได้ จึงออกแบบให้บันทึก audit trail ทุกขั้นตอน",
         cover="pool-01", cover_cap="แดชบอร์ด Billing สรุปปริมาณและราคาก๊าซรายเดือน",
         metrics=[("3", ["แหล่งจัดหาก๊าซ", "LNG · Gulf · Myanmar"]),
                  ("100%", ["ทุกการคำนวณ", "บันทึก audit trail"])],
         gallery=[]),

    dict(key="verso", eyebrow="Aquila Studio · Procurement",
         title=["Verso PO/PR", "Procurement Workflow"],
         desc="ระบบจัดซื้อตั้งแต่ยื่นคำขอจนถึงรับสินค้า รองรับ Multi-Level Approval, "
              "Budget Validation และ Purchase Order Automation",
         platform="Web Application (Internal)",
         work="Backend · Approval Workflow · Budget Control",
         stack="TypeScript · PostgreSQL · Docker",
         overview="ออกแบบ flow อนุมัติหลายระดับตามโครงสร้างองค์กร ตรวจงบก่อนใช้จริงเพื่อกัน budget overrun "
                  "และมี Audit Trail สำหรับการตรวจสอบภายใน",
         cover="verso-01", cover_cap="หน้าแรกของระบบจัดซื้อ พร้อมรายการคำขอเบิกจ่าย",
         gal_lede="หน้า Purchase Order ที่ออกจากคำขอซื้อ พร้อมตารางรายการ สรุปภาษี และช่องอนุมัติ",
         gallery=[("verso-02", "หน้า Purchase Order พร้อมตารางรายการและสรุปภาษี")]),

    dict(key="prolab", eyebrow="Aquila Studio · HealthTech",
         title=["Prolab", "Health Service Platform"],
         desc="Web Application สำหรับจัดการบริการตรวจสุขภาพและห้องปฏิบัติการ ครอบคลุม Patient Information, "
              "Package Management, Online Appointment และ Booking Lifecycle",
         platform="Responsive Web · Customer & Back-office",
         work="Backend · REST API · Lab Result Flow",
         stack="TypeScript · PostgreSQL · HL7 · Docker",
         overview="ร่วมพัฒนาฝั่ง backend ของ AI Health Analytics ที่ช่วยประเมินความเสี่ยงจากผลตรวจเป็นข้อมูลประกอบให้แพทย์ "
                  "และปรับ flow ของ Lab Result System จนทีมลดเวลารอผลจาก 3–7 วัน เหลือ 1–3 วัน",
         link="ai.prolab.co.th",
         cover="prolab-01", cover_cap="หน้าแรกของ ProLab Web Application",
         metrics=[("3–7 → 1–3", ["วัน — เวลารอผลแล็บ", "หลังทีมปรับ flow ระบบ"])],
         gal_lede="ฝั่งผู้ใช้ — นัดหมาย แพ็กเกจตรวจสุขภาพ และหน้าเว็บไซต์หลัก",
         gallery=[("prolab-02", "หน้ารายการนัดหมายของผู้ใช้"),
                  ("prolab-03", "แพ็กเกจตรวจสุขภาพบนหน้าเว็บ"),
                  ("prolab-04", "หน้าเว็บไซต์หลักของระบบ")]),

    dict(key="movesook", eyebrow="Side project · SaaS Marketplace",
         title=["MoveSook", "On-demand Moving Platform"],
         desc="Two-sided marketplace เชื่อมผู้ใช้กับคนขับขนย้ายผ่าน LINE/LIFF พร้อมหลังบ้านสำหรับทีม Admin",
         platform="LINE Mini App · Responsive Web · Admin",
         work="ทำเองทั้งหมด — Architecture · Backend · Frontend · Deploy",
         stack="Turborepo · Next.js · Hono · Prisma · PostgreSQL · Zod · LINE LIFF",
         overview="ออกแบบเป็น Monorepo แยก 3 แอป (web · admin · api) ที่ใช้ end-to-end type-safe RPC ตั้งแต่ DB ถึง UI "
                  "มี RBAC 4 ระดับ (USER / DRIVER / ADMIN / SYSTEM) และ Job State Machine สำหรับ business logic การจับคู่งาน",
         link="movesook.com",
         cover="movesook-01", cover_cap="หน้าแรก เรียกรถรับจ้างขนย้าย",
         metrics=[("3", ["แอปใน Monorepo", "web · admin · api"]),
                  ("4", ["ระดับสิทธิ์ RBAC", "USER · DRIVER · ADMIN · SYSTEM"])],
         gal_lede="หน้าอธิบายบริการและพื้นที่ให้บริการ — ออกแบบให้ผู้ใช้เข้าใจขั้นตอนก่อนโพสต์งานขนย้าย",
         gallery=[("movesook-02", "จุดเด่นของบริการ และขั้นตอนใช้งาน 3 ขั้นตอน"),
                  ("movesook-03", "พื้นที่ให้บริการ และคำถามที่พบบ่อย")]),

    dict(key="icmt", eyebrow="Aquila Studio · EdTech",
         title=["ICMT", "ระบบบริการจัดการแท็บเล็ต"],
         desc="แพลตฟอร์ม Monorepo สำหรับบริหารจัดการแท็บเล็ตที่แจกจ่ายให้โรงเรียน ครอบคลุมทะเบียนทรัพย์สิน "
              "ยืม–คืน ตรวจรับคืน ตั๋วแจ้งเคลมซ่อม และรายงานรายโรงเรียน",
         platform="Web · Admin Console · Background Worker",
         work="Backend · RBAC · Queue · Audit Log",
         stack="Turborepo · Next.js · Hono · Prisma · PostgreSQL · Redis",
         overview="แยกเป็น Web, Admin, Storage Service และ Background Worker รองรับ Role-based Access "
                  "(SYSTEM / USER / ADMIN), JWT + Refresh Token Rotation, Redis Queue สำหรับงานเบื้องหลัง "
                  "และ Audit Log ที่ตรวจสอบย้อนหลังได้",
         cover="icmt-01", cover_cap="แดชบอร์ดภาพรวมสถานะอุปกรณ์ พร้อมรายงานรายโรงเรียน",
         metrics=[("2,000+", ["อุปกรณ์ในระบบ", "กระจายหลายโรงเรียน"]),
                  ("4", ["บริการแยกกัน", "Web · Admin · Storage · Worker"])],
         gal_lede="ฝั่งผู้ใช้ทั่วไป — หน้าแรกของระบบ และหน้ายืมอุปกรณ์พร้อมสถานะการยืม",
         gallery=[("icmt-02", "หน้าแรกของระบบสำหรับผู้ใช้ทั่วไป"),
                  ("icmt-03", "หน้าผู้ใช้ สำหรับยืมอุปกรณ์และดูสถานะการยืม")]),

    dict(key="mmert", eyebrow="Side project · Emergency Medical",
         title=["M-MERT", "Maritime Emergency Medical"],
         desc="ระบบสั่งการกู้ภัยทางการแพทย์ฉุกเฉินทางน้ำ รองรับภารกิจ Water / Land / Air ตั้งแต่รับแจ้งเหตุ "
              "จัดทีมอาสาสมัคร ติดตามรถพยาบาลแบบเรียลไทม์ ไปจนถึงประสานงานกับโรงพยาบาลปลายทาง",
         platform="Web · Command Dashboard · Field Units",
         work="ทำเองทั้งหมด — Architecture · Backend · Frontend · Deploy",
         stack="Bun · Hono 4 · Effect · Next.js 16 · PostgreSQL 16 · Prisma 6 · SSE",
         overview="สั่งการผ่าน SSE push — คำสั่งจากแพทย์ถึงหน่วยหน้างานทันทีโดยไม่ต้องรอ refresh มี Triage 4 ระดับ, "
                  "Vital Signs Monitoring, Waypoint Navigation (GPS / UTM / MGRS) และ Medical Order System แบบเรียลไทม์",
         cover="mmert-01", cover_cap="แดชบอร์ดศูนย์สั่งการ รวมสถานะภารกิจและทรัพยากรในหน้าเดียว",
         metrics=[("9", ["โมดูลหลัก", "Mission · Triage · Vital · Order"]),
                  ("4", ["ระดับ Triage", "แดง · เหลือง · เขียว · ดำ"])],
         gal_lede="ฝั่งปฏิบัติการ — ติดตามยานพาหนะเรียลไทม์ จัดการผู้ป่วยตามระดับ Triage ติดตามสถานการณ์ภารกิจ และประสานโรงพยาบาล",
         gallery=[("mmert-02", "แผนที่ติดตามยานพาหนะแบบเรียลไทม์"),
                  ("mmert-03", "ระบบจัดการผู้ป่วยพร้อมระดับ Triage"),
                  ("mmert-04", "หน้าติดตามสถานการณ์ของภารกิจ"),
                  ("mmert-05", "ระบบจัดการโรงพยาบาลเชื่อมต่อ")]),

    dict(key="clinic-rag", eyebrow="Side project · AI · RAG",
         title=["Clinic Booking", "+ AI Chatbot"],
         desc="ต่อ AI Chatbot แบบ RAG เข้ากับระบบจองคิวคลินิก — ผู้ป่วยถามเรื่องบริการ ราคา และเวลาเปิดทำการได้ผ่าน LINE "
              "โดยดึงคำตอบจากฐานข้อมูลของคลินิกผ่าน Vector Search",
         platform="LINE LIFF · Web Admin · Worker",
         work="ทำเองทั้งหมด — RAG Pipeline · Backend · Admin",
         stack="Turborepo · Next.js · LangChain · OpenAI SDK · Qdrant · Prisma · BullMQ · MinIO",
         overview="RAG Pipeline ด้วย LangChain + OpenAI + Qdrant Vector Database, LINE LIFF ให้คนไข้จองคิวผ่านมือถือ "
                  "โดยไม่ต้องลงแอป, Admin Dashboard สำหรับเจ้าหน้าที่จัดการนัดหมาย และ Worker Service สำหรับงานพื้นหลัง",
         cover="clinic-rag-01", cover_cap="Analytics Dashboard วิเคราะห์คุณภาพการตอบของแชท AI",
         gal_lede="ฝั่งหลังบ้าน — จัดการบริการและแพ็กเกจของคลินิก และข้อมูลแพทย์ (ปิดชื่อและรูปเพื่อความเป็นส่วนตัว)",
         gallery=[("clinic-rag-02", "หน้าจัดการบริการและแพ็กเกจของคลินิก"),
                  ("clinic-rag-03", "หน้าจัดการข้อมูลแพทย์")]),
]

INDEX_ITEMS = [
    ("01", "ERC", "erc-01", "Energy Platform · Aquila Studio"),
    ("02", "Pool Manager", "pool-01", "Gas Management · Aquila Studio"),
    ("03", "Verso PO/PR", "verso-01", "Procurement · Aquila Studio"),
    ("04", "Prolab", "prolab-01", "HealthTech · Aquila Studio"),
    ("05", "MoveSook", "movesook-01", "SaaS Marketplace · Side project"),
    ("06", "ICMT", "icmt-01", "EdTech · Aquila Studio"),
    ("07", "M-MERT", "mmert-01", "Emergency Medical · Side project"),
    ("08", "Clinic Booking + AI", "clinic-rag-01", "AI · RAG · Side project"),
]

TIMELINE = [
    ("กรกฎาคม 2565 — ปัจจุบัน", "Mid-level Backend Developer",
     "บริษัท อะควิลา สตูดิโอ จำกัด · HealthTech · Energy · Enterprise"),
    ("ตุลาคม 2564 — มิถุนายน 2565", "Frontend Developer",
     "บริษัท คิวบ์ออฟไนน์ จำกัด · Angular · Responsive UI"),
    ("ตุลาคม 2563 — ตุลาคม 2564", "Fullstack Developer",
     "บริษัท เดต้าฮอลิค จำกัด · POS · CRM · ระบบภายในองค์กร"),
]

PILLARS = [
    ("01 · Backend Development", "clinic-erp",
     "สถาปัตยกรรม Microservices Monorepo ของ Clinic ERP",
     "ออกแบบและพัฒนา RESTful API กับฐานข้อมูล ดูแลตั้งแต่คุย requirement วางโครงสร้างข้อมูล จนขึ้น production",
     "TypeScript · Node.js · Bun · NestJS · Hono · PostgreSQL · Prisma"),
    ("02 · AI & Integration", "meawsook",
     "AI Vector Search Pipeline ของ MeawSook",
     "ต่อ AI แบบ RAG เข้ากับระบบเดิม ทำ Vector Search และเชื่อมระบบภายนอก เช่น HL7, LINE LIFF และ Payment",
     "OpenAI SDK · LangChain · Qdrant · pgvector · HL7 Parser"),
    ("03 · DevOps & Delivery", "mmert",
     "สถาปัตยกรรม Real-time Event-driven ของ M-MERT",
     "ดูแล deployment ด้วย Docker และ CI/CD บน cloud จัดการงาน background ด้วย queue และงาน real-time ด้วย SSE",
     "Docker · GitHub Actions · Nginx · Linux · Redis · BullMQ · SSE"),
]

SIDE = [
    ("SaaS ERP · LINE Mini App", "StockSook",
     "ERP + POS ขนาดเล็กสำหรับร้านค้า SME — สต๊อก ขายหน้าร้าน คำนวณต้นทุนเฉลี่ยถ่วงน้ำหนัก และรายงาน เข้าใช้ผ่าน LINE",
     "Turborepo · Next.js 15 · Hono · PostgreSQL · LINE LIFF"),
    ("B2B SaaS · Multi-tenant", "Clinic ERP",
     "ระบบ ERP + CRM สำหรับคลินิก แยก Workspace รายสาขา 6 แอปใน Monorepo ใช้ fp-ts จัดการ error",
     "Hono 4 · fp-ts · SuperTokens · BullMQ · Cloudflare R2"),
    ("Social Good · AI Matching", "MeawSook",
     "LINE Mini App ตามหาแมวหายและบริจาคอาหารแมวแบบตรวจสอบย้อนหลังได้ ใช้ Multimodal Embedding จับคู่ใบหน้าแมว",
     "Voyage Multimodal-3 · pgvector · Hono · Cloudflare Workers"),
    ("Clinical Ops · Internal Tool", "EMS-ECI",
     "ระบบบันทึก Checklist ประจำวันและทะเบียนผู้ป่วยแผลกดทับสำหรับพยาบาล พร้อม Excel Report สำหรับประชุม QI",
     "Next.js 16 · React 19 · Prisma 7 · better-auth · ExcelJS"),
]

SKILLS = [
    ("Backend", "Node.js · TypeScript · Bun · NestJS · ExpressJS · Fastify · HonoJS · ElysiaJS · Python · FastAPI · Go · C# .NET · REST API · tRPC · GraphQL"),
    ("Frontend", "HTML · CSS · JavaScript · React · Angular · Next.js · Tailwind CSS · Bootstrap · PrimeNG · PrimeReact · Apollo Client"),
    ("Database & ORM", "MySQL · PostgreSQL · MSSQL · SQLite · MongoDB · Redis · Qdrant · pgvector · Prisma · TypeORM · Sequelize"),
    ("DevOps & Tools", "Docker · Docker Compose · Nginx · Git · GitHub Actions · AWS (EC2/S3/RDS) · Linux · Turborepo · pnpm · n8n"),
    ("AI & Integration", "OpenAI SDK · LangChain · RAG · HL7 Parser · Playwright · BullMQ · Zod · NextAuth · better-auth · JWT · Pino · ExcelJS · Sharp · LINE LIFF"),
    ("Architecture & Concepts", "OOP · MVC · System Design · API Security & Authentication · Microservices · Multi-tenant SaaS · Event-driven / SSE · Responsive Design"),
]


# ---------------------------------------------------------------- เตรียมรูป
def prepare_images():
    """แปลง .webp เป็น .jpg เพราะ PowerPoint บางเวอร์ชันไม่รองรับ webp"""
    os.makedirs(BUILD, exist_ok=True)
    made = 0
    for name in sorted(os.listdir(ASSETS)):
        if not name.endswith(".webp"):
            continue
        dst = os.path.join(BUILD, name[:-5] + ".jpg")
        if not os.path.exists(dst):
            subprocess.run(["magick", os.path.join(ASSETS, name), "-strip",
                            "-quality", "88", dst], check=True)
            made += 1
    return made


def img(name):
    p = os.path.join(BUILD, name + ".jpg")
    if not os.path.exists(p):
        p = os.path.join(BUILD, name + ".png")
    if not os.path.exists(p):
        raise SystemExit("ไม่พบรูป %s ใน %s" % (name, BUILD))
    return p


# ---------------------------------------------------------------- สไลด์
def slide_cover(prs):
    s = new_slide(prs)
    ff = s.shapes.build_freeform(emu(1920), emu(0))
    ff.add_line_segments([(emu(1920), emu(1080)), (emu(154), emu(0))], close=True)
    sh = ff.convert_to_shape()
    sh.fill.solid(); sh.fill.fore_color.rgb = WASH
    sh.line.fill.background(); sh.shadow.inherit = False

    text(s, 80, 690, 900, 40, "Selected work", size=24, color=ACCENT, bold=True,
         caps=True, letter=3.2)
    text(s, 80, 736, 1100, 280,
         [[("Muhammadhisam", {})], [("Pala", {})]],
         size=128, color=NAVY, bold=True, spacing=0.98, fit=False)
    text(s, 80, 972, 1100, 60, "Software Developer Portfolio 2026", size=46, color=NAVY)
    text(s, 80, 1032, 1100, 40, "มูฮัมหมัดฮีซาม ปาล๊ะ (ซัง) — Backend / Fullstack",
         size=26, color=RGBColor(0x7C, 0x88, 0x94))

    text(s, 1000, 84, 840, 130,
         [[("muhammadhisam.pala@gmail.com", {"bold": True, "color": NAVY})],
          [("090-163-0867 · LINE: hisam023", {})],
          [("github.com/muhammmadhisam", {})]],
         size=19, color=RGBColor(0x6E, 0x7A, 0x87), align=PP_ALIGN.RIGHT, spacing=1.6)
    return s


def slide_index(prs):
    s = new_slide(prs)
    chrome(s, "Contents", "Selected work",
           "แปดระบบที่ผมมีส่วนร่วมพัฒนาและได้ใช้งานจริง — แต่ละตัวมีหน้าอธิบายและภาพหน้าจอในหน้าถัดไป")
    for i, (no, nm, im, tag) in enumerate(INDEX_ITEMS):
        col, row = i % 4, i // 4
        x, y = 80 + col * 440, 286 + row * 350
        rect(s, x, y, 410, 256, fill=PANEL, line=HAIR)
        picture(s, x, y, 410, 256, img(im))
        text(s, x, y + 268, 410, 22, no, size=13, color=ACCENT, bold=True, letter=2.4)
        text(s, x, y + 290, 410, 34, nm, size=22, color=NAVY, bold=True)
        text(s, x, y + 322, 410, 26, tag, size=15, color=LABEL)
    return s


def slide_expertise(prs):
    s = new_slide(prs)
    rect(s, 760, 0, 1160, 1080, fill=WASH)
    chrome(s, "About me", "Expertise", None, narrow_foot=True)
    text(s, 80, 196, 600, 160,
         "Backend Developer สาย TypeScript / Node.js / Bun ทำงานตั้งแต่คุย requirement "
         "ออกแบบฐานข้อมูลและ API เขียนโค้ด ไปจนถึง deploy และดูแลระบบหลังขึ้น production",
         size=24, color=BODY, spacing=1.45)

    label(s, 80, 430, 600, "จุดที่ถนัด")
    text(s, 80, 458, 600, 90,
         "ออกแบบ REST API และฐานข้อมูลให้ทีมหน้าบ้านต่อง่าย จัดการงาน queue / background job "
         "และดูแล deployment ด้วย Docker + CI/CD", size=19, color=BODY, spacing=1.5)
    label(s, 80, 570, 600, "กำลังเรียนรู้")
    text(s, 80, 598, 600, 90,
         "สาย AI / RAG — ต่อ Vector Search และ LLM เข้ากับระบบที่มีอยู่ เพื่อลดงานซ้ำ ๆ ของผู้ใช้",
         size=19, color=BODY, spacing=1.5)

    for i, (num, cap) in enumerate([("5+", "ปีในสายพัฒนา"), ("10+", "ระบบที่ใช้งานจริง"),
                                    ("3", "บริษัทที่ร่วมงาน")]):
        x = 80 + i * 200
        text(s, x, 800, 190, 70, num, size=52, color=NAVY, bold=True, spacing=1.0, fit=False)
        text(s, x, 862, 190, 30, cap, size=16, color=LABEL)

    label(s, 860, 150, 900, "Career timeline")
    y = 200
    for when, role, where in TIMELINE:
        rect(s, 860, y, 980, 1, fill=HAIR)
        text(s, 860, y + 26, 980, 26, when, size=16, color=LABEL, letter=1.3)
        text(s, 860, y + 56, 980, 44, role, size=32, color=NAVY, bold=True)
        text(s, 860, y + 102, 980, 34, where, size=20, color=BODY)
        y += 172
    rect(s, 860, y, 980, 1, fill=HAIR)
    text(s, 860, y + 40, 980, 90,
         "5+ ปีที่ผ่านมาส่วนใหญ่อยู่กับงานสาย HealthTech, Energy และระบบภายในองค์กร "
         "เคยทำทั้ง Microservices, multi-tenant SaaS และงาน queue / background job",
         size=18, color=LABEL, spacing=1.5)
    return s


def slide_what_i_do(prs):
    s = new_slide(prs)
    chrome(s, "About me", "What I do",
           "สามเรื่องหลักที่ผมรับผิดชอบในทุกโปรเจกต์ — ตั้งแต่ออกแบบสถาปัตยกรรม จนขึ้น production และดูแลต่อ")
    for i, (name, dia, cap, desc, stack) in enumerate(PILLARS):
        x = 80 + i * 600
        rect(s, x, 300, 560, 300, fill=WASH, line=HAIR)
        picture_fit(s, x + 16, 316, 528, 268, img(dia))
        text(s, x, 612, 560, 26, cap, size=14, color=LABEL)
        label(s, x, 652, 560, name)
        text(s, x, 684, 560, 140, desc, size=18, color=BODY, spacing=1.5)
        text(s, x, 852, 560, 60, stack, size=16, color=LABEL, spacing=1.4)
    return s


def slide_project(prs, p):
    s = new_slide(prs)
    rect(s, 680, 0, 1240, 1080, fill=WASH)
    rect(s, 870, 160, 860, 700, fill=PANEL, line=HAIR)
    picture(s, 870, 160, 860, 700, img(p["cover"]))
    text(s, 870, 880, 860, 40, p["cover_cap"], size=16, color=LABEL)
    if p.get("metrics"):
        metrics(s, 870, 912, p["metrics"])

    text(s, 80, 22, 600, 30, p["eyebrow"], size=15, color=LABEL, bold=True, caps=True, letter=2.1)
    rect(s, 80, 62, 600, 4, fill=ACCENT)
    text(s, 80, 100, 570, 140, [[(t, {})] for t in p["title"]],
         size=46, color=NAVY, bold=True, spacing=1.15, fit=False)
    text(s, 80, 268, 560, 180, p["desc"], size=21, color=BODY, spacing=1.55)

    y = 492
    for lab, val in (("Platform", p["platform"]), ("Work done", p["work"]),
                     ("Stack", p["stack"])):
        label(s, 80, y, 560, lab)
        text(s, 80, y + 26, 560, 60, val, size=19, color=BODY, spacing=1.45)
        y += 102
    label(s, 80, y, 560, "Overview")
    text(s, 80, y + 26, 560, 180, p["overview"], size=19, color=BODY, spacing=1.5)

    if p.get("link"):
        text(s, 80, 952, 560, 34, p["link"] + "  ↗", size=19, color=NAVY, bold=True)
    foot(s, narrow=True)
    return s


GAL_LAYOUT = {
    1: [(360, 1200, 620)],
    2: [(80, 860, 535), (980, 860, 535)],
    3: [(80, 560, 470), (680, 560, 470), (1280, 560, 470)],
    4: [(80, 410, 345), (530, 410, 345), (980, 410, 345), (1430, 410, 345)],
}


def slide_gallery(prs, p):
    shots = p["gallery"]
    n = len(shots)
    if not n:
        return None
    s = new_slide(prs)
    eyebrow = " ".join(p["title"])
    chrome(s, eyebrow, "More screens", p.get("gal_lede", ""))
    for (x, w, h), (f, cap) in zip(GAL_LAYOUT[n], shots):
        rect(s, x, 300, w, h, fill=PANEL, line=HAIR)
        picture(s, x, 300, w, h, img(f))
        text(s, x, 320 + h, w, 50, cap, size=15, color=LABEL, spacing=1.35)
    return s


def slide_side(prs):
    s = new_slide(prs)
    chrome(s, "Outside of work", "More side projects",
           "โปรเจกต์ที่ทำนอกเวลางาน เพื่อฝึกออกแบบระบบและลองเทคโนโลยีใหม่ ๆ บางตัวเปิดใช้จริงแล้ว บางตัวยังพัฒนาอยู่")
    for i, (tag, name, desc, stack) in enumerate(SIDE):
        col, row = i % 2, i // 2
        x, y = 80 + col * 900, 330 + row * 300
        rect(s, x, y, 840, 1, fill=HAIR)
        text(s, x, y + 24, 840, 26, tag, size=15, color=LABEL, bold=True, caps=True, letter=2.0)
        text(s, x, y + 56, 840, 44, name, size=28, color=NAVY, bold=True)
        text(s, x, y + 106, 840, 100, desc, size=18, color=BODY, spacing=1.5)
        text(s, x, y + 218, 840, 34, stack, size=16, color=LABEL)
    return s


def slide_skills(prs):
    s = new_slide(prs)
    chrome(s, "Toolbox", "Skills & technologies", "เครื่องมือที่ใช้สร้างงานในแต่ละวัน")
    for i, (name, items) in enumerate(SKILLS):
        col, row = i % 3, i // 3
        x, y = 80 + col * 600, 330 + row * 330
        label(s, x, y, 560, name)
        text(s, x, y + 30, 560, 240, items, size=19, color=BODY, spacing=1.55)
    return s


def slide_thanks(prs):
    s = new_slide(prs)
    ff = s.shapes.build_freeform(emu(1920), emu(0))
    ff.add_line_segments([(emu(1920), emu(1080)), (emu(154), emu(0))], close=True)
    sh = ff.convert_to_shape()
    sh.fill.solid(); sh.fill.fore_color.rgb = WASH
    sh.line.fill.background(); sh.shadow.inherit = False

    text(s, 80, 700, 900, 40, "Let's talk", size=24, color=ACCENT, bold=True, caps=True, letter=3.2)
    text(s, 80, 748, 1200, 170, "Thank you", size=128, color=NAVY, bold=True, spacing=0.98, fit=False)
    text(s, 80, 950, 1300, 60,
         "muhammadhisam.pala@gmail.com  |  090-163-0867", size=46, color=NAVY)
    text(s, 80, 1016, 1500, 40,
         "LINE: hisam023 · linkedin.com/in/muhammadhisam-pala-45b83825a · github.com/muhammmadhisam",
         size=24, color=RGBColor(0x7C, 0x88, 0x94))
    text(s, 1000, 84, 840, 80,
         [[("Additional work available at", {"bold": True, "color": NAVY})],
          [("github.com/muhammmadhisam/profile-muhammadhisam", {})]],
         size=19, color=RGBColor(0x6E, 0x7A, 0x87), align=PP_ALIGN.RIGHT, spacing=1.6)
    return s


# ---------------------------------------------------------------- main
def main():
    if not os.path.isdir(ASSETS):
        raise SystemExit("ไม่พบโฟลเดอร์ %s" % ASSETS)
    prepare_images()

    prs = Presentation()
    prs.slide_width = Emu(int(SLIDE_W_IN * 914400))
    prs.slide_height = Emu(int(SLIDE_H_IN * 914400))

    names = []
    slide_cover(prs); names.append("Cover — Muhammadhisam Pala")
    slide_index(prs); names.append("Selected work")
    slide_expertise(prs); names.append("Expertise & career timeline")
    slide_what_i_do(prs); names.append("What I do")
    for p in PROJECTS:
        slide_project(prs, p); names.append(" ".join(p["title"]))
        if slide_gallery(prs, p) is not None:
            names.append(" ".join(p["title"]) + " — More screens")
    slide_side(prs); names.append("More side projects")
    slide_skills(prs); names.append("Skills & technologies")
    slide_thanks(prs); names.append("Thank you")

    for slide, name in zip(prs.slides, names):
        slide.notes_slide.notes_text_frame.text = name

    prs.core_properties.title = "Muhammadhisam Pala — Software Developer Portfolio 2026"
    prs.core_properties.author = "Muhammadhisam Pala"
    prs.core_properties.subject = "Software Developer Portfolio"
    prs.save(OUT)
    print("%d สไลด์ -> %s" % (len(names), OUT))


if __name__ == "__main__":
    sys.exit(main())
