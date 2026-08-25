#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""แปลง Muhammadhisam-Pala-Portfolio-2026.pdf เป็นไฟล์ .pptx (สไลด์ละ 1 ภาพเต็มหน้า)

ต้องมี:
    pip install python-pptx
    brew install poppler        # ให้ได้คำสั่ง pdftoppm
    brew install imagemagick    # ให้ได้คำสั่ง magick

วิธีใช้ (สร้าง PDF ใหม่ก่อนถ้าแก้ portfolio-deck.html):
    python3 -m http.server 4180
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome" \
        --headless=new --no-pdf-header-footer \
        --print-to-pdf=Muhammadhisam-Pala-Portfolio-2026.pdf \
        http://localhost:4180/portfolio-deck.html
    python3 make-pptx.py
"""
import glob
import os
import shutil
import subprocess
import tempfile

from pptx import Presentation
from pptx.util import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
PDF = os.path.join(HERE, "Muhammadhisam-Pala-Portfolio-2026.pdf")
OUT = os.path.join(HERE, "Muhammadhisam-Pala-Portfolio-2026.pptx")

# ชื่อสไลด์ ใช้เป็นชื่อ shape และ speaker notes ให้เปิดใน outline แล้วอ่านออก
TITLES = [
    "Cover — Muhammadhisam Pala",
    "Selected work",
    "Expertise & career timeline",
    "What I do",
    "ERC — Centralized Energy Data Platform",
    "ERC — More screens",
    "Pool Manager — Gas Pool Management",
    "Verso PO/PR — Procurement Workflow",
    "Verso PO/PR — More screens",
    "Prolab — Health Service Platform",
    "Prolab — More screens",
    "MoveSook — On-demand Moving Platform",
    "MoveSook — More screens",
    "ICMT — ระบบบริการจัดการแท็บเล็ต",
    "ICMT — More screens",
    "M-MERT — Maritime Emergency Medical",
    "M-MERT — More screens",
    "Clinic Booking + AI Chatbot",
    "Clinic Booking + AI — More screens",
    "More side projects",
    "Skills & technologies",
    "Thank you",
]


def render_pages(workdir):
    """แตกทุกหน้าของ PDF เป็น JPEG ความละเอียด 2200px"""
    subprocess.run(
        ["pdftoppm", "-png", "-r", "110", PDF, os.path.join(workdir, "s")],
        check=True,
    )
    pages = []
    for png in sorted(glob.glob(os.path.join(workdir, "s-*.png"))):
        jpg = png[:-4] + ".jpg"
        subprocess.run(["magick", png, "-strip", "-quality", "88", jpg], check=True)
        os.remove(png)
        pages.append(jpg)
    return pages


def build(pages):
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for i, img in enumerate(pages):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(img, 0, 0, width=prs.slide_width, height=prs.slide_height)
        title = TITLES[i] if i < len(TITLES) else "Slide %d" % (i + 1)
        slide.shapes[0].name = title
        slide.notes_slide.notes_text_frame.text = title

    prs.core_properties.title = "Muhammadhisam Pala — Software Developer Portfolio 2026"
    prs.core_properties.author = "Muhammadhisam Pala"
    prs.core_properties.subject = "Software Developer Portfolio"
    prs.save(OUT)
    return len(pages)


def main():
    if not os.path.exists(PDF):
        raise SystemExit("ไม่พบ %s — สร้าง PDF ก่อน (ดูคำอธิบายด้านบนของไฟล์นี้)" % PDF)
    workdir = tempfile.mkdtemp(prefix="deck-")
    try:
        pages = render_pages(workdir)
        if not pages:
            raise SystemExit("แตกหน้า PDF ไม่ได้")
        print("%d สไลด์ -> %s" % (build(pages), OUT))
    finally:
        shutil.rmtree(workdir, ignore_errors=True)


if __name__ == "__main__":
    main()
